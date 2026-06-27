"""
Export all analyzed inspections from the local database to Excel.
Run from the fda_483_workflow directory:
    python3 export_to_excel.py
"""
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

from db import init_db, get_session
from db.models import Inspection

# ── Colours ──────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
CRITICAL_FILL = PatternFill("solid", fgColor="FF0000")
MAJOR_FILL    = PatternFill("solid", fgColor="FF6600")
MINOR_FILL    = PatternFill("solid", fgColor="FFD700")
ALT_FILL      = PatternFill("solid", fgColor="EBF3FB")
HEADER_FONT   = Font(bold=True, color="FFFFFF", size=11)
BOLD          = Font(bold=True)

THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RISK_FILL = {"Critical": CRITICAL_FILL, "Major": MAJOR_FILL, "Minor": MINOR_FILL}


def _cell(ws, row, col, value, font=None, fill=None, wrap=False, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    if font:
        c.font = font
    elif bold:
        c.font = BOLD
    if fill:
        c.fill = fill
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    c.border = BORDER
    return c


def build_summary_sheet(wb, rows):
    ws = wb.active
    ws.title = "All Inspections"

    headers = [
        "Month", "Firm Name", "Country", "City/State",
        "Inspection Date", "# Observations", "Risk Level",
        "Executive Summary", "Key Themes", "Top GMP Gaps", "Recommendations",
    ]

    for col, h in enumerate(headers, 1):
        _cell(ws, 1, col, h, font=HEADER_FONT, fill=HEADER_FILL)
    ws.row_dimensions[1].height = 22

    for row_idx, r in enumerate(rows, 2):
        alt = ALT_FILL if row_idx % 2 == 0 else None
        risk = r["risk_level"] or ""
        risk_fill = RISK_FILL.get(risk, alt)

        date = r["inspection_end_date"]
        month = date.strftime("%Y-%m") if date else ""
        location = ", ".join(filter(None, [r["city"], r["state"], r["country"]]))

        themes = "; ".join(r["key_themes"] or [])
        gaps = "\n".join(
            f"[{g.get('framework','')}] {g.get('reference','')} — {g.get('gap_description','')}"
            for g in (r["top_gmp_gaps"] or [])
        )

        values = [
            month, r["firm_name"] or "", r["country"] or "", location,
            str(date) if date else "", r["num_observations"] or 0, risk,
            r["executive_summary"] or "", themes, gaps, r["recommendations"] or "",
        ]

        for col, val in enumerate(values, 1):
            fill = risk_fill if col == 7 else alt
            _cell(ws, row_idx, col, val, fill=fill, wrap=(col >= 8))

        ws.row_dimensions[row_idx].height = max(60, 15 * max(1, len(r["executive_summary"] or "") // 80))

    widths = [10, 32, 12, 18, 14, 8, 10, 55, 35, 45, 45]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def build_monthly_sheets(wb, rows):
    from collections import defaultdict, Counter
    by_month = defaultdict(list)
    for r in rows:
        key = r["inspection_end_date"].strftime("%Y-%m") if r["inspection_end_date"] else "Unknown"
        by_month[key].append(r)

    for month in sorted(by_month):
        ws = wb.create_sheet(title=month)
        month_rows = by_month[month]

        ws.merge_cells("A1:G1")
        c = ws.cell(row=1, column=1, value=f"FDA CGMP Warning Letters — {month}  ({len(month_rows)} inspections)")
        c.font = Font(bold=True, size=13, color="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        risks = Counter(r["risk_level"] or "Unknown" for r in month_rows)
        ws.cell(row=2, column=1, value=f"Critical: {risks.get('Critical',0)}  |  Major: {risks.get('Major',0)}  |  Minor: {risks.get('Minor',0)}")
        ws.cell(row=2, column=1).font = Font(bold=True, color="444444")
        ws.merge_cells("A2:G2")
        ws.row_dimensions[2].height = 18

        headers = ["Firm", "Country", "Date", "Observations", "Risk", "Key Themes", "Recommendations"]
        for col, h in enumerate(headers, 1):
            _cell(ws, 3, col, h, font=HEADER_FONT, fill=HEADER_FILL)
        ws.row_dimensions[3].height = 20

        for i, r in enumerate(month_rows, 4):
            risk = r["risk_level"] or ""
            alt = ALT_FILL if i % 2 == 0 else None
            themes = "; ".join(r["key_themes"] or [])

            row_vals = [
                r["firm_name"] or "", r["country"] or "",
                str(r["inspection_end_date"]) if r["inspection_end_date"] else "",
                r["num_observations"] or 0, risk, themes, r["recommendations"] or "",
            ]
            for col, val in enumerate(row_vals, 1):
                fill = RISK_FILL.get(risk, alt) if col == 5 else alt
                _cell(ws, i, col, val, fill=fill, wrap=(col >= 6))
            ws.row_dimensions[i].height = 45

        widths = [32, 12, 12, 10, 10, 40, 50]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w

        ws.freeze_panes = "A4"


def main():
    init_db()
    with get_session() as session:
        from sqlalchemy.orm import joinedload
        inspections = (
            session.query(Inspection)
            .join(Inspection.analysis)
            .options(
                joinedload(Inspection.analysis),
                joinedload(Inspection.observations),
            )
            .order_by(Inspection.inspection_end_date)
            .all()
        )
        # Detach from session by converting to plain dicts
        rows = []
        for insp in inspections:
            a = insp.analysis
            rows.append({
                "firm_name": insp.firm_name,
                "city": insp.city,
                "state": insp.state,
                "country": insp.country,
                "inspection_end_date": insp.inspection_end_date,
                "num_observations": insp.num_observations,
                "risk_level": a.risk_level if a else "",
                "executive_summary": a.executive_summary if a else "",
                "key_themes": a.key_themes if a else [],
                "top_gmp_gaps": a.top_gmp_gaps if a else [],
                "recommendations": a.recommendations if a else "",
            })

    print(f"Found {len(rows)} analyzed inspections")

    if not rows:
        print("No data yet — run the backfill first.")
        return

    wb = openpyxl.Workbook()
    build_summary_sheet(wb, rows)
    build_monthly_sheets(wb, rows)

    out = Path(__file__).parent / "data" / "reports" / f"FDA_CGMP_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    wb.save(out)
    print(f"\nExcel report saved: {out}")
    print("Open it in Finder with:  open \"" + str(out) + "\"")


if __name__ == "__main__":
    main()
