#!/usr/bin/env python3
"""
Convert a single WeChat conversation exported by WechatExporter into one Word (.docx) file.

This is the last-mile step of the export pipeline (see README.md):

    unencrypted iPhone backup  ->  WechatExporter (HTML/Text)  ->  THIS SCRIPT  ->  chat.docx

WechatExporter can output Text/HTML/PDF but not .docx, so this script fills that gap.
It runs entirely offline: no network, no cloud service, no LLM. The only hard
dependency is python-docx.

Usage:
    python3 html_to_docx.py INPUT.html OUTPUT.docx
    python3 html_to_docx.py INPUT.txt  OUTPUT.docx      # plain-text export also works

Input type is chosen by file extension:
    .html / .htm  -> tags are stripped, each message block becomes a paragraph
    .txt          -> each non-empty line becomes a paragraph

Tip: if you have `pandoc` installed, `pandoc INPUT.html -o OUTPUT.docx` is an
even simpler one-liner. This script exists for users without pandoc and for a
little more control over the output.
"""

import sys
from html.parser import HTMLParser


class _TextBlockParser(HTMLParser):
    """Extract readable text from HTML, inserting a paragraph break at block boundaries.

    Uses only the Python standard library so the script has no dependency beyond
    python-docx. It is intentionally forgiving: it does not rely on any specific
    WechatExporter CSS class, so it keeps working across template versions.
    """

    _BLOCK_TAGS = {
        "div", "p", "br", "tr", "li", "table",
        "h1", "h2", "h3", "h4", "h5", "h6",
    }
    _SKIP_TAGS = {"script", "style", "head"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self._current = []
        self._skip_depth = 0

    def _flush(self):
        text = " ".join(self._current).strip()
        # collapse runs of whitespace
        text = " ".join(text.split())
        if text:
            self.blocks.append(text)
        self._current = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif tag in self._BLOCK_TAGS:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self._SKIP_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in self._BLOCK_TAGS:
            self._flush()

    def handle_data(self, data):
        if self._skip_depth == 0 and data.strip():
            self._current.append(data)

    def close(self):
        super().close()
        self._flush()


def extract_paragraphs(path):
    """Return a list of message paragraphs from an .html/.htm or .txt export."""
    lower = path.lower()
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        content = fh.read()

    if lower.endswith((".html", ".htm")):
        parser = _TextBlockParser()
        parser.feed(content)
        parser.close()
        return parser.blocks

    if lower.endswith(".txt"):
        return [line.strip() for line in content.splitlines() if line.strip()]

    raise SystemExit(
        f"Unsupported input extension for {path!r}. "
        "Expected .html, .htm, or .txt (WechatExporter output)."
    )


def write_docx(paragraphs, out_path, title="WeChat Conversation"):
    try:
        from docx import Document
    except ImportError:
        raise SystemExit(
            "python-docx is not installed. Run:\n"
            "    pip3 install -r requirements.txt\n"
            "  (or)  pip3 install python-docx"
        )

    doc = Document()
    doc.add_heading(title, level=0)
    if not paragraphs:
        doc.add_paragraph("(No messages were found in the export.)")
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(out_path)


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        raise SystemExit("Usage: python3 html_to_docx.py INPUT.(html|txt) OUTPUT.docx")

    in_path, out_path = argv[1], argv[2]
    paragraphs = extract_paragraphs(in_path)
    write_docx(paragraphs, out_path)
    print(f"Wrote {out_path} ({len(paragraphs)} message paragraph(s)) from {in_path}")


if __name__ == "__main__":
    main(sys.argv)
